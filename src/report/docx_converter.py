"""Conversor Markdown -> DOCX do relatório científico (Roadmap V15.4 / Spec G8).

Preserva estrutura de seções, tabelas e listas usando ``python-docx``.
"""

from __future__ import annotations

import re
from pathlib import Path

from src.report.base_converter import ORIGIN_ANCHOR_RE, ReportConverter, split_marks, strip_markup_helpers

_HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
_TABLE_ROW_RE = re.compile(r"^\s*\|(.+)\|\s*$")
_TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")
_BULLET_RE = re.compile(r"^\s*[-*]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^\s*\d+\.\s+(.*)$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _add_marked_runs(paragraph, text: str, bold: bool = False) -> None:
    """Acrescenta ``text`` ao parágrafo, com as marcas do relatório (origem sobrescrita, não verificado em realce)."""
    from docx.enum.text import WD_COLOR_INDEX

    for mark in split_marks(ORIGIN_ANCHOR_RE.sub("", text)):
        if mark.kind == "text":
            if mark.value:
                paragraph.add_run(mark.value).bold = bold or None
        elif mark.kind == "origin":
            paragraph.add_run(mark.value).font.superscript = True
        else:
            run = paragraph.add_run(mark.value)
            run.bold = True
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _fill_runs(paragraph, text: str) -> None:
    pos = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > pos:
            _add_marked_runs(paragraph, text[pos:match.start()])
        _add_marked_runs(paragraph, match.group(1), bold=True)
        pos = match.end()
    if pos < len(text):
        _add_marked_runs(paragraph, text[pos:])


def _add_paragraph_with_bold(document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    _fill_runs(paragraph, text)


class DOCXConverter(ReportConverter):
    """Converte ``relatorio_final.md`` para um documento Word (.docx)."""

    def convert(self, markdown_path: Path, output_path: Path) -> None:
        from docx import Document

        markdown_text = strip_markup_helpers(Path(markdown_path).read_text(encoding="utf-8"))
        lines = markdown_text.splitlines()

        document = Document()
        i = 0
        while i < len(lines):
            line = lines[i]

            if line.strip().startswith("```"):
                i += 1
                code_lines = []
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                i += 1
                code_paragraph = document.add_paragraph("\n".join(code_lines))
                code_paragraph.style = document.styles["Normal"]
                for run in code_paragraph.runs:
                    run.font.name = "Courier New"
                continue

            if _TABLE_ROW_RE.match(line) and i + 1 < len(lines) and _TABLE_SEPARATOR_RE.match(lines[i + 1]):
                table_lines = [line]
                i += 2
                while i < len(lines) and _TABLE_ROW_RE.match(lines[i]):
                    table_lines.append(lines[i])
                    i += 1
                rows = [
                    [cell.strip() for cell in row.strip().strip("|").split("|")]
                    for row in table_lines
                ]
                num_cols = len(rows[0])
                table = document.add_table(rows=len(rows), cols=num_cols)
                table.style = "Light Grid Accent 1"
                for r, row in enumerate(rows):
                    cells = row + [""] * (num_cols - len(row))
                    for c in range(num_cols):
                        # Runs (não ``cell.text``): as marcas de origem e de não verificado valem dentro de tabelas.
                        cell_paragraph = table.cell(r, c).paragraphs[0]
                        _fill_runs(cell_paragraph, cells[c])
                continue

            heading_match = _HEADING_RE.match(line)
            if heading_match:
                level = len(heading_match.group(1))
                document.add_heading(heading_match.group(2), level=level)
                i += 1
                continue

            bullet_match = _BULLET_RE.match(line)
            if bullet_match:
                _add_paragraph_with_bold(document, bullet_match.group(1), style="List Bullet")
                i += 1
                continue

            numbered_match = _NUMBERED_RE.match(line)
            if numbered_match:
                _add_paragraph_with_bold(document, numbered_match.group(1), style="List Number")
                i += 1
                continue

            if not line.strip():
                i += 1
                continue

            _add_paragraph_with_bold(document, line)
            i += 1

        document.save(str(output_path))
