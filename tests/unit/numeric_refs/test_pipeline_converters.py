"""Pipeline do relatório, catálogo, conversores, contagens e contratos dos agentes (spec numeric-provenance)."""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest

from src.numeric_refs import provenance_file
from src.numeric_refs.catalog import build_reference_catalog
from src.numeric_refs.literal_metrics import write_findings
from src.report.base_converter import ReportConverterFactory
from src.report.pipeline import SOURCE_FILE, ReportPipeline, ReportState, build_default_pipeline, delimit_section

from .helpers import make_resolver, new_ledger, record_run

pytestmark = pytest.mark.unit


@pytest.fixture
def session(tmp_path):
    outputs = tmp_path / "outputs"
    ledger = new_ledger(outputs)
    return outputs, ledger


def _state(outputs, ledger, text, **kw) -> ReportState:
    session_dir = outputs / "s1"
    session_dir.mkdir(parents=True, exist_ok=True)
    return ReportState(
        session_id="s1", session_dir=session_dir, source_text=text, text=text,
        resolver=make_resolver(outputs, ledger, **kw), counts={"subtarefa": 1},
    )


def test_pipeline_ponta_a_ponta(session):
    outputs, ledger = session
    e1 = record_run(ledger, outputs, metrics={"rmse": 0.4498, "acc": 0.87345}, unidades={"rmse": "mm"})
    source = (
        "# Relatório\n\n## Resultados\n"
        f"O RMSE foi {{{{res:{e1}/rmse}}}} e a acurácia {{{{res:{e1}/acc}}}}; houve 12 falhas.\n"
        "Foi executada uma subtarefa.\n\n"
        "<!-- execution-metadata:begin -->\ntokens 1234\n<!-- execution-metadata:end -->\n"
    )
    state = _state(outputs, ledger, source)
    target = build_default_pipeline().build(state)
    assert target.name == "relatorio_final.md"
    final = target.read_text(encoding="utf-8")
    assert "0,4498 mm [R1]" in final and "0,87345 [R2]" in final
    assert "12 [não verificado] falhas" in final and "tokens 1234" in final
    assert final.count("[não verificado]") == 1
    # apêndice delimitado, com âncoras
    assert "<!-- numeric-provenance:begin -->" in final and "## Origem dos números" in final
    assert '<a id="origem-R1"></a>R1' in final and e1 in final
    payload = json.loads((outputs / "s1" / "proveniencia_numerica.json").read_text(encoding="utf-8"))
    assert payload["contagens"] == {"res": 2, "calc": 0, "src": 0, "nao_verificados": 1, "metricas_literais": 0}
    assert [c["codigo"] for c in payload["codigos"]] == ["R1", "R2"]
    assert not list((outputs / "s1").glob("*.tmp"))  # gravação atômica


def test_apendice_completo_com_res_calc_e_src(session):
    """Scenario: Apêndice completo."""
    outputs, ledger = session
    e1 = record_run(ledger, outputs, metrics={"a": 0.5, "b": 0.25})
    record_search = __import__("src.numeric_refs.sources", fromlist=["x"]).record_search_sources
    record_search(outputs / "s1", "q", [{"url": "https://e.org/p", "title": "P", "snippet": "valor de 4 amostras"}])
    text = (
        f"{{{{res:{e1}/a}}}} {{{{res:{e1}/b}}}} {{{{calc:res:{e1}/a / res:{e1}/b}}}} "
        "{{src:https://e.org/p#valor de 4 amostras}}"
    )
    state = _state(outputs, ledger, text, session_ids=["s1"])
    pipeline = build_default_pipeline()
    state = pipeline.run(state)
    appendix = "\n".join(state.sections)
    for code in ("R1", "R2", "C1", "S1"):
        assert f"origem-{code}" in appendix
    assert "execução registrada" in appendix and "cálculo determinístico" in appendix and "fonte citada" in appendix


def test_etapa_extra_e_secao_extra_registradas_por_outras_mudancas(session):
    outputs, ledger = session
    pipeline = build_default_pipeline()
    order: list[str] = []

    def extra(state):
        order.append("claims")
        state.text += "\nafirmação 12"
        return state

    pipeline.register_stage("claims", extra, after="resolve")
    pipeline.register_section("claims", lambda s: "tabela de afirmações 9", "Afirmações")
    assert pipeline.stage_names == ["resolve", "claims", "verify"]
    with pytest.raises(ValueError):
        pipeline.register_stage("claims", extra)
    with pytest.raises(ValueError):
        pipeline.register_stage("x", extra, after="inexistente")
    state = pipeline.run(_state(outputs, ledger, "texto"))
    assert order == ["claims"] and "12 [não verificado]" in state.text  # a etapa roda antes da verificação
    assert any("<!-- claims:begin -->" in s and "tabela de afirmações 9" in s for s in state.sections)


def test_literais_aparecem_no_apendice_e_na_contagem(session):
    outputs, ledger = session
    e = record_run(ledger, outputs, metrics={"acc": 0.95})
    from src.numeric_refs.literal_metrics import read_literal_metrics, scan_code

    findings = scan_code('save_experiment_artifacts("t", {}, {"acc": 0.95})')
    write_findings(outputs / "s1" / "treinar", e, "h" * 64, findings)

    state = _state(outputs, ledger, f"acc {{{{res:{e}/acc}}}}", literal_metrics=read_literal_metrics(outputs / "s1"))
    build_default_pipeline().build(state)
    final = (outputs / "s1" / "relatorio_final.md").read_text(encoding="utf-8")
    assert "0,95 [R1] [literal no código]" in final and "Métricas gravadas como literais" in final
    assert provenance_file.NumericProvenanceReader(outputs).read("s1")["metricas_literais"] == 1


def test_leitor_do_grupo_proveniencia(session):
    """Scenario: Leitura do grupo proveniência."""
    outputs, ledger = session
    e = record_run(ledger, outputs, metrics={"a": 0.5, "b": 0.25, "c": 0.1})
    text = (
        f"{{{{res:{e}/a}}}} {{{{res:{e}/b}}}} {{{{res:{e}/c}}}} {{{{calc:res:{e}/a * 2}}}} "
        "{{src:x#y 1}} e 7 e 8"
    )
    build_default_pipeline().build(_state(outputs, ledger, text))
    outputs_session = outputs / "s1"
    payload = json.loads((outputs_session / "proveniencia_numerica.json").read_text(encoding="utf-8"))
    payload["metricas_literais"] = [{"exec_id": e, "metrica": "a", "linha": 1, "padrao": "P1"}]
    payload["contagens"]["metricas_literais"] = 1
    (outputs_session / "proveniencia_numerica.json").write_text(json.dumps(payload), encoding="utf-8")
    result = provenance_file.NumericProvenanceReader(outputs).read("s1")
    assert result == {
        "numeros_por_origem": {"res": 3, "calc": 1, "src": 1},
        "numeros_nao_verificados": 3,  # 7, 8 e a citação {{src}} que não resolve (listada como motivo)
        "metricas_literais": 1,
    }


def test_catalogo_lista_referencias_prontas_para_copiar(session):
    outputs, ledger = session
    e = record_run(ledger, outputs, metrics={"rmse": 0.4498, "n": 3}, unidades={"rmse": "mm"}, parameters={"k": 3})
    data = json.loads((outputs / "s1" / "treinar" / "metrics.json").read_text())
    data["parameters"] = {"k": 3}
    text = build_reference_catalog(
        session_id="s1", session_dir=outputs / "s1", ledger=ledger, metrics_by_task={"treinar": data}
    )
    assert f"{{{{res:{e}/rmse}}}} = 0.4498 mm" in text and f"{{{{res:{e}/n}}}} = 3" in text
    assert f"{{{{res:{e}/param.k}}}} = 3" in text and "escreva as referências abaixo, não os números" in text


def test_catalogo_sem_execucao_registrada(session):
    outputs, ledger = session
    text = build_reference_catalog(
        session_id="s1", session_dir=outputs / "s1", ledger=ledger,
        metrics_by_task={"orfa": {"metrics": {"a": 1}}},
    )
    assert "sem execução registrada" in text and "{{res:" not in text
    ledger.store.unavailable = True
    assert "indisponível" in build_reference_catalog(
        session_id="s1", session_dir=outputs / "s1", ledger=ledger, metrics_by_task={}
    )


# --- relatório: marcadores do orquestrador --------------------------------------------------------------------------

def test_secoes_do_orquestrador_sao_delimitadas_no_relatorio():
    from src.report.report_model import ReportMetadata, build_report_data, parse_narrative, render_report_markdown

    data = build_report_data(
        title="T", request="p", metrics_by_task={"a": {"metrics": {"acc": 0.9}}},
        metadata=ReportMetadata(duration_s=12.0, tokens_in=100, tokens_out=50, cost_usd=0.0123),
    )
    narrative = parse_narrative(json.dumps({
        "resumo_executivo": "r", "contexto_e_objetivo": "c", "metodologia": "m", "analise_divergencias": "a",
        "limitacoes": "l", "proximos_passos": "p", "confianca_nivel": "alto", "confianca_justificativa": "j",
    }))
    markdown = render_report_markdown(data, narrative)
    for name in ("report-results", "researcher-decisions", "execution-metadata"):
        assert f"<!-- {name}:begin -->" in markdown and f"<!-- {name}:end -->" in markdown
    from src.numeric_refs.verifier import verify_numbers

    assert verify_numbers(markdown).nao_verificados == []  # nenhum número do orquestrador é marcado


# --- conversores ---------------------------------------------------------------------------------------------------

REPORT = (
    "# Relatório\n\n"
    "A acurácia foi 0,95 [não verificado] e o RMSE 0,87 [R1] [literal no código].\n\n"
    "| Métrica | Valor |\n|---|---|\n| acc | 0,95 [não verificado] |\n| rmse | 0,87 [R1] |\n\n"
    "<!-- numeric-provenance:begin -->\n## Origem dos números\n\n"
    "| Código | Valor exato |\n|---|---|\n| <a id=\"origem-R1\"></a>R1 | 0.87 |\n"
    "<!-- numeric-provenance:end -->\n"
)


def test_conversor_html(tmp_path):
    """Scenario: HTML."""
    source = tmp_path / "relatorio_final.md"
    source.write_text(REPORT, encoding="utf-8")
    out = tmp_path / "r.html"
    ReportConverterFactory.create("html").convert(source, out)
    html = out.read_text(encoding="utf-8")
    assert '<mark class="nao-verificado">não verificado</mark>' in html
    assert '<mark class="literal">literal no código</mark>' in html
    assert '<sup class="origem"><a href="#origem-R1">R1</a></sup>' in html
    assert 'id="origem-R1"' in html  # o alvo do link existe no apêndice
    assert "numeric-provenance" not in html  # os delimitadores de seção não aparecem


def test_conversor_html_nao_toca_codigo(tmp_path):
    source = tmp_path / "r.md"
    source.write_text("texto `[R1]` e\n\n```\n[não verificado]\n```\n", encoding="utf-8")
    out = tmp_path / "r.html"
    ReportConverterFactory.create("html").convert(source, out)
    html = out.read_text(encoding="utf-8")
    assert "<code>[R1]</code>" in html and "<mark" not in html


def test_conversor_docx_inclusive_em_tabela(tmp_path):
    """Scenario: DOCX."""
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    source = tmp_path / "r.md"
    source.write_text(REPORT, encoding="utf-8")
    out = tmp_path / "r.docx"
    ReportConverterFactory.create("docx").convert(source, out)
    document = Document(str(out))

    def runs_of(paragraphs):
        return [run for p in paragraphs for run in p.runs]

    body_runs = runs_of(document.paragraphs)
    marked = [r for r in body_runs if r.text == "não verificado"]
    assert marked and all(r.bold and r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in marked)
    assert any(r.text == "R1" and r.font.superscript for r in body_runs)
    cells = [cell for t in document.tables for row in t.rows for cell in row.cells]
    table_runs = [run for cell in cells for run in runs_of(cell.paragraphs)]
    assert any(r.text == "não verificado" and r.bold for r in table_runs)
    assert not any("<!--" in p.text or "<a id" in p.text for p in document.paragraphs)


def test_conversor_latex(tmp_path):
    """Scenario: LaTeX."""
    source = tmp_path / "r.md"
    source.write_text(REPORT, encoding="utf-8")
    out = tmp_path / "r.tex"
    ReportConverterFactory.create("latex").convert(source, out)
    tex = out.read_text(encoding="utf-8")
    assert r"\usepackage{xcolor}" in tex and r"\colorbox{yellow}{\textbf{não verificado}}" in tex
    assert r"\textsuperscript{\hyperlink{origem-R1}{R1}}" in tex and r"\hypertarget{origem-R1}{}" in tex
    assert r"\colorbox{yellow}{\textbf{literal no código}}" in tex and "<!--" not in tex
    assert re.search(r"\\textsuperscript\{\\hyperlink\{origem-R1\}\{R1\}\}", tex)


# --- contratos dos agentes -----------------------------------------------------------------------------------------

def test_instrucao_do_summarizer_descreve_as_referencias():
    """Scenario: Instrução do Summarizer."""
    from agents.summarizer.agent import AGENT_INSTRUCTION

    for token in ("{{res:", "{{calc:", "{{src:"):
        assert token in AGENT_INSTRUCTION
    assert "calcule a divergência percentual" not in AGENT_INSTRUCTION.lower()
    assert "Não calcule divergências" in " ".join(AGENT_INSTRUCTION.split())


def test_instrucoes_do_developer_e_do_curator():
    from agents.curator.agent import AGENT_INSTRUCTION as CURATOR
    from agents.developer.agent import AGENT_INSTRUCTION as DEVELOPER

    assert "unidades=" in DEVELOPER and "literal" in DEVELOPER
    assert "{{res:" in CURATOR and "RECUSAM referências malformadas" in CURATOR


def test_orquestrador_acrescenta_o_pipeline_a_sintese(tmp_path):
    from tests.unit.test_report_model import _loop, _synth, narrative_json

    loop, _orch = _loop(tmp_path, [narrative_json()])
    __import__("asyncio").run(_synth(loop))
    session_dir = Path(_orch.output_manager.base_dir) / "sess"
    assert (session_dir / SOURCE_FILE).is_file() and (session_dir / "relatorio_final.md").is_file()
    assert (session_dir / "proveniencia_numerica.json").is_file()
    assert "## Origem dos números" in (session_dir / "relatorio_final.md").read_text(encoding="utf-8")


def test_secao_delimitada_helper():
    assert delimit_section("x", "Título", "corpo") == "<!-- x:begin -->\n## Título\ncorpo\n<!-- x:end -->"
    assert isinstance(ReportPipeline().section_names, list)
