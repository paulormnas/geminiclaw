"""Testes unitários dos conversores de formato do relatório científico
(Roadmap V15.4 / Spec G8).
"""

import shutil
import subprocess
from pathlib import Path

import pytest

from src.report.base_converter import ReportConverterFactory
from src.report.docx_converter import DOCXConverter
from src.report.html_converter import HTMLConverter
from src.report.latex_converter import LaTeXConverter

SAMPLE_MARKDOWN = """# Relatório de Teste

## Resumo Executivo
Este é um **resumo** com *ênfase* e `código inline`, incluindo 100% de acurácia & precisão.

## Resultados

| Subtarefa | Acurácia | Seed |
|---|---|---|
| treinar_modelo | 0.87 | 42 |
| validar | 0.91 | 42 |

## Próximos Passos
- Item um
- Item dois com **negrito**

1. Primeiro passo
2. Segundo passo

```python
print("hello world")
```
"""


@pytest.mark.unit
class TestReportConverterFactory:
    def test_cria_latex_converter(self) -> None:
        assert isinstance(ReportConverterFactory.create("latex"), LaTeXConverter)

    def test_cria_html_converter(self) -> None:
        assert isinstance(ReportConverterFactory.create("html"), HTMLConverter)

    def test_cria_docx_converter(self) -> None:
        assert isinstance(ReportConverterFactory.create("docx"), DOCXConverter)

    def test_formato_invalido_lanca_erro_com_lista_de_formatos(self) -> None:
        with pytest.raises(ValueError, match="latex, html, docx"):
            ReportConverterFactory.create("pdf")

    def test_case_insensitive(self) -> None:
        assert isinstance(ReportConverterFactory.create("LaTeX"), LaTeXConverter)


@pytest.mark.unit
class TestLaTeXConverter:
    def test_produz_arquivo_nao_vazio(self, tmp_path: Path) -> None:
        md_path = tmp_path / "relatorio_final.md"
        md_path.write_text(SAMPLE_MARKDOWN, encoding="utf-8")
        out_path = tmp_path / "relatorio_final.tex"

        LaTeXConverter().convert(md_path, out_path)

        assert out_path.exists()
        content = out_path.read_text(encoding="utf-8")
        assert len(content) > 0
        assert r"\documentclass{article}" in content
        assert r"\begin{document}" in content
        assert r"\end{document}" in content

    def test_titulo_h1_nao_duplicado_como_secao(self, tmp_path: Path) -> None:
        md_path = tmp_path / "r.md"
        md_path.write_text(SAMPLE_MARKDOWN, encoding="utf-8")
        out_path = tmp_path / "r.tex"
        LaTeXConverter().convert(md_path, out_path)

        content = out_path.read_text(encoding="utf-8")
        assert r"\title{Relatório de Teste}" in content
        assert r"\section{Relatório de Teste}" not in content

    def test_tabela_convertida_para_tabular(self, tmp_path: Path) -> None:
        md_path = tmp_path / "r.md"
        md_path.write_text(SAMPLE_MARKDOWN, encoding="utf-8")
        out_path = tmp_path / "r.tex"
        LaTeXConverter().convert(md_path, out_path)

        content = out_path.read_text(encoding="utf-8")
        assert r"\begin{tabular}" in content
        assert "treinar\\_modelo & 0.87 & 42" in content

    def test_escapa_caracteres_especiais(self, tmp_path: Path) -> None:
        md_path = tmp_path / "r.md"
        md_path.write_text(SAMPLE_MARKDOWN, encoding="utf-8")
        out_path = tmp_path / "r.tex"
        LaTeXConverter().convert(md_path, out_path)

        content = out_path.read_text(encoding="utf-8")
        assert r"100\%" in content
        assert r"\&" in content

    @pytest.mark.skipif(shutil.which("pdflatex") is None, reason="pdflatex não instalado neste ambiente")
    def test_compila_com_pdflatex(self, tmp_path: Path) -> None:
        """Validação end-to-end: o .tex gerado deve ser compilável com pdflatex."""
        md_path = tmp_path / "relatorio_final.md"
        md_path.write_text(SAMPLE_MARKDOWN, encoding="utf-8")
        tex_path = tmp_path / "relatorio_final.tex"
        LaTeXConverter().convert(md_path, tex_path)

        result = subprocess.run(
            ["pdflatex", "-interaction=nonstopmode", "-halt-on-error", tex_path.name],
            cwd=tmp_path,
            capture_output=True,
            text=True,
            timeout=60,
        )
        assert result.returncode == 0, result.stdout[-2000:]
        assert (tmp_path / "relatorio_final.pdf").exists()


@pytest.mark.unit
class TestHTMLConverter:
    def test_produz_arquivo_nao_vazio(self, tmp_path: Path) -> None:
        md_path = tmp_path / "relatorio_final.md"
        md_path.write_text(SAMPLE_MARKDOWN, encoding="utf-8")
        out_path = tmp_path / "relatorio_final.html"

        HTMLConverter().convert(md_path, out_path)

        assert out_path.exists()
        content = out_path.read_text(encoding="utf-8")
        assert "<!DOCTYPE html>" in content
        assert "<table>" in content
        assert "@media print" in content
        assert "@media (prefers-color-scheme: dark)" in content

    def test_titulo_extraido_do_h1(self, tmp_path: Path) -> None:
        md_path = tmp_path / "r.md"
        md_path.write_text(SAMPLE_MARKDOWN, encoding="utf-8")
        out_path = tmp_path / "r.html"
        HTMLConverter().convert(md_path, out_path)

        content = out_path.read_text(encoding="utf-8")
        assert "<title>Relatório de Teste</title>" in content


@pytest.mark.unit
class TestDOCXConverter:
    def test_produz_arquivo_docx_valido(self, tmp_path: Path) -> None:
        from docx import Document

        md_path = tmp_path / "relatorio_final.md"
        md_path.write_text(SAMPLE_MARKDOWN, encoding="utf-8")
        out_path = tmp_path / "relatorio_final.docx"

        DOCXConverter().convert(md_path, out_path)

        assert out_path.exists()
        document = Document(str(out_path))
        assert len(document.paragraphs) > 0
        assert len(document.tables) == 1
        assert document.tables[0].cell(0, 0).text == "Subtarefa"
